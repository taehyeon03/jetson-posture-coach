from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR

OUT = Path(__file__).parent
S = {
 'P1': ('Lee et al., Applied Sciences 2023, 13(19), 10935', 'https://doi.org/10.3390/app131910935'),
 'P2': ('Yu et al., Lite-HRNet, CVPR 2021, pp. 10440–10450', 'https://openaccess.thecvf.com/content/CVPR2021/html/Yu_Lite-HRNet_A_Lightweight_High-Resolution_Network_CVPR_2021_paper.html'),
 'P3': ('Wang et al., Lite Pose, CVPR 2022, pp. 13126–13136', 'https://openaccess.thecvf.com/content/CVPR2022/html/Wang_Lite_Pose_Efficient_Architecture_Design_for_2D_Human_Pose_Estimation_CVPR_2022_paper.html'),
 'B1': ('WHO, 신체활동 부족 보도자료, 2024-06-26', 'https://www.who.int/news/item/26-06-2024-nearly-1.8-billion-adults-at-risk-of-disease-from-not-doing-enough-physical-activity'),
 'B2': ('Hinge Health, AI care tools, 2025-10-21', 'https://ir.hingehealth.com/news/news-details/2025/Hinge-Health-unveils-new-AI-powered-care-tools/default.aspx'),
 'B3': ('Hinge Health R&D, 기기 지연 최적화, 2026-08-04', 'https://rnd.hingehealth.com/blog/cutting-latency-by-up-to-15-percent'),
 'M1': ('EverEx 공식 뉴스룸, 2025년 국내 사업 동향', 'https://www.everex.kr/ko/news?year=2025'),
 'M2': ('EverEx 공식 제품·서비스 소개, 2026-09-13 열람', 'https://www.everex.kr/'),
 'M3': ('Hinge Health, 2025 연간 실적, 2026-02-10', 'https://ir.hingehealth.com/news/news-details/2026/Hinge-Health-reports-fourth-quarter-and-full-year-2025-financial-results/default.aspx'),
 'H1': ('NVIDIA, Jetson Nano 공식 사양', 'https://developer.nvidia.com/embedded/jetson-nano'),
 'H2': ('NVIDIA, JetPack 4 최종 릴리스 공지, 2024-11-22', 'https://forums.developer.nvidia.com/t/announcing-end-of-life-for-nvidia-jetpack-4-with-the-release-of-jetpack-4-6-6/314300'),
}

ABSTRACT = ('본 연구는 Jetson Nano 4GB와 단일 RGB 카메라를 활용하여 컴퓨터 작업 중 전방머리자세와 상체 기울어짐을 감지하고 자세 교정을 유도하는 온디바이스 시스템을 제안한다. '
 '경량 딥러닝 자세 추정 모델로 측면 영상의 귀·어깨·엉덩이 좌표를 추출하고, 개인별 기준 자세에 대한 상대적 변위와 각도 변화를 분석한다. '
 '키포인트 신뢰도와 자세 지속 시간을 함께 고려하여 일시적인 움직임에 따른 불필요한 알림을 줄이고, 기준 이탈이 지속될 때 시각·청각 피드백을 제공한다. '
 '모델 학습은 외부 PC에서 수행하며, 장치에서는 단일 사용자 추론과 간단한 후처리를 실행하여 4GB 메모리 제약에 대응한다. '
 '원본 영상은 기본적으로 저장하거나 외부로 전송하지 않고 자세 지표와 알림 이력을 로컬에 기록한다. '
 '평가는 사용자 단위로 분리한 데이터에서 자세 분류 성능, 시간당 오경보, 처리 지연, 최대 메모리 사용량 및 알림 후 자세 복귀율을 측정한다. '
 '이를 통해 제한된 엣지 자원에서 개인화된 자세 감지와 지속 가능한 피드백의 구현 가능성을 검증하고자 한다.')

slides = [
 dict(section='01 · COVER', title='Jetson Nano 4GB 기반\n개인화 자세 교정 지원 시스템', sub='경량 딥러닝으로 전방머리자세와 상체 기울어짐 감지', cards=[('학과 · 학번 · 이름', '[학과 입력]  /  [학번 입력]  /  [이름 입력]'),('팀 · 과목', '[팀명 입력]  /  [과목명 입력]')], takeaway='연구 제안서 · 문헌 조사 기준일 2026.09.13', cover=True),
 dict(section='02 · TEAM', title='팀 소개 및 역할', sub='팀 인원이 미정인 상태의 역할 배분안 — 실제 팀원에 맞춰 통합·조정', cards=[('팀원 A · [이름 / 학번]', '기획·문헌 조사\n연구 질문, 배경·시장 조사, 발표 자료 총괄'),('팀원 B · [이름 / 학번]', 'Computer Vision\n데이터 수집·라벨링, 경량 자세 추정, 개인화 판단'),('팀원 C · [이름 / 학번]', 'Computer Science\nNano 배포, 메모리·지연 최적화, 로컬 기록'),('팀원 D · [이름 / 학번]', '평가·사용자 경험\n실험 설계, 오경보·사용성 평가, 피드백 화면')], takeaway='공동 산출물: 동작 시제품 + 사용자 분리 평가 + 재현 가능한 실행 설정'),
 dict(section='03 · CONTENTS', title='목차', sub='과제의 1–6항목을 본문에 반영하고, 논문·초록은 부록으로 구성', cards=[('01–03 · 시작', '표지\n팀 소개 및 역할\n목차'),('04 · 관련 주제', '대분류 / 중분류 / 소분류\n대상 사용자와 적용 환경'),('05 · Background', '연구 배경\n2025–2026 기사·트렌드'),('06 · Motivation', '국내 시장성 및 도입 필요성\n해외 시장성 및 구매자 관점')], takeaway='부록: CV·CS 기여 / 4GB 구현 계획 / 논문 3편 / 검증 계획 / 초록 / 참고문헌'),
 dict(section='04 · TOPIC', title='디지털 헬스 / 일상 자세 모니터링 / 개인화 교정 지원', sub='단일 사용자가 앉아 컴퓨터를 사용하는 환경부터 시작', cards=[('대분류 · 디지털 헬스', '일상 속 건강관리와 자기관리 지원'),('중분류 · 비접촉 자세 모니터링', '측면 RGB 카메라로 머리·상체 위치 변화 관찰'),('소분류 · 엣지 AI 자세 교정 지원', 'Jetson Nano 4GB에서 전방머리자세와\n상체 기울어짐을 감지하고 알림 제공')], takeaway='적용 범위: 자세 변화 인지와 행동 유도. 임상 진단이나 치료 효과는 별도 검증 대상.'),
 dict(section='05 · BACKGROUND 1 / 2', title='일상에서 자세 변화를 알아차리게 하는 도구', sub='문제 인식 → 생활 속 관찰 → 적절한 시점의 피드백', cards=[('사회적 배경', 'WHO: 2022년 전 세계 성인 31%가\n권장 신체활동량에 미달\n2024년 발표된 생활 건강관리 관련 지표 [B1]'),('연구의 출발점', '컴퓨터 작업 중 머리·몸통의 자세를\n영상으로 분류한 선행연구가 존재 [P1]\n본 과제는 일상 사용 중의 피드백에 주목'),('Inspiration', '자세를 계속 의식하도록 요구하기보다\n기준 자세에서 벗어난 시간이 길어질 때\n사용자가 알아차릴 수 있도록 알림')], takeaway='신체활동 부족 통계는 거북목 유병률이나 자세 알림의 건강 효과를 의미하지 않는다.', sources=['B1','P1']),
 dict(section='05 · BACKGROUND 2 / 2', title='최신 동향: 카메라 기반 평가와 온디바이스 피드백', sub='2025–2026 공개 자료에서 확인한 서비스·개발 방향', cards=[('2025.10 · 서비스 확대', 'Hinge Health가 Movement Analysis 발표\n카메라 기반 관절 각도·대칭성 등 평가를\n실제 근골격계 관리 서비스에 연결 [B2]'),('2026.08 · 현장 성능 개선', 'Hinge Health 개발팀은 사용자 기기에서\n동작하는 CV 처리의 지연 최적화를 소개\n기기 성능 차이 대응이 개발 과제 [B3]'),('본 과제에 주는 시사점', '인식 정확도와 함께 반응성·지속 사용 고려\n고정된 저사양 장치에서 성능 검증\n개인별 기준과 알림 빈도 제어를 결합')], takeaway='기업 자료는 상용화 방향의 근거이며, 본 시스템의 성능·효과를 증명하는 자료는 아니다.', sources=['B2','B3']),
 dict(section='06 · MOTIVATION 1 / 2', title='국내 시장성: 산업보건·재활에서 일상 관리로', sub='기술 구조 대신 고객, 도입 경로, 구매 이유를 조사', cards=[('확인된 국내 움직임', 'EverEx 공식 뉴스룸: 2025.02\n대한산업보건협회와 산재 예방 협력 MOU\n2025.03에는 검사 키오스크 관련 보도 [M1]'),('현재 공급과 고객', 'EverEx는 동작 평가와 재활 솔루션 제공 [M2]\n인접 수요: 의료·재활, 근로자 건강관리\n기존 제품의 존재는 사업 기회의 근거'),('우리 팀의 시장 가설', '초기 고객: 대학 연구실·소규모 사무실\n구매 가치: 간단한 설치, 착용 부담 감소,\n영상 외부 전송 없는 개인 자세 관리')], takeaway='국내 자세 알림 기기의 시장 규모·지불의사는 확인되지 않음 → 인터뷰와 시범 도입으로 검증.', sources=['M1','M2']),
 dict(section='06 · MOTIVATION 2 / 2', title='해외 시장성: 근골격계 디지털 관리의 지불 수요', sub='2026년 발표된 2025 연간 실적을 활용한 인접시장 조사', cards=[('매출 · 성장', 'Hinge Health 2025년 매출\n5억 8,790만 달러\n전년 대비 51% 증가 [M3]'),('고객 · 이용 규모', '2025년 말 기업 고객 2,830곳\n회원 782,890명 [M3]\n조직 단위 구매 수요를 보여주는 사례'),('도입 필요성 · 검증 항목', '예상 구매자: 고용주·대학·개인 사용자\n확인할 가치: 설치비·유지비·알림 만족도\n유료 전환 의사와 반복 사용률을 조사')], takeaway='위 수치는 한 기업의 근골격계 관리 사업 실적이다. 거북목 교정기 시장 규모로 환산하지 않는다.', sources=['M3']),
 dict(section='APPENDIX · CONTRIBUTION', title='CV와 CS에 기여할 연구 질문', sub='“4GB 장치에서 정확도·오경보·지연의 균형을 어떻게 확보할 것인가?”', cards=[('CV · 개인차와 일시적 움직임', '귀–어깨 상대 위치와 몸통 기울기를 분석\n개인 기준 대비 변화 + 키포인트 신뢰도\n지속 시간 조건으로 오경보 감소 여부 검증'),('CS · 제한된 자원에서 지속 실행', '고정 길이 프레임 버퍼와 추론 주기 제어\n모델 정밀도·입력 크기별 지연과 RAM 비교\n원본 영상 대신 로컬 자세 통계 관리'),('기여를 입증하는 비교', '고정 임계값 vs 개인 기준 적용\n단일 프레임 vs 지속 시간 적용\nFP32 vs FP16: 정확도·지연·메모리 비교')], takeaway='제안한 기여는 연구 가설이며, 기존 방법보다 우수하다는 주장은 실험 이후 판단한다.'),
 dict(section='APPENDIX · 4GB DESIGN', title='Jetson Nano 4GB에서 실행하기 위한 설계', sub='학습은 외부 PC, Nano는 단일 사용자 추론·후처리 담당 [H1]', cards=[('입력·모델 범위', '측면 카메라 1대 / 사용자 1명 / 고정 ROI\nLite-HRNet-18을 1차 후보로 검토\n입력 256×192, batch 1에서 시작'),('메모리 관리', 'TensorRT FP16 변환 가능성 우선 확인\n모델 1개만 상주, 프레임 버퍼 길이 제한\nOS·공유 GPU·UI 포함 전체 RAM 계측'),('운용·성공 기준', '설계 목표: 5–10 FPS, 최대 RAM 3.5GB 이하\n30분 연속 실행 시 OOM·swap 의존 없음\nJetPack 4.6.6 호환성 사전 점검 [H2]')], takeaway='목표치는 실측 결과가 아니다. 모델 크기만으로 실행 메모리나 FPS를 보장할 수 없다.', sources=['H1','H2','P2']),
 dict(section='APPENDIX · PAPER 01', title='컴퓨터 작업 중 목·척추 자세 분류', sub='Classifying Poor Postures of the Neck and Spine in Computer Work\nby Using Image and Skeleton Analysis', cards=[('Paper information', 'Jaeeun Lee, Hongseok Choi,\nKyeongmin Yum, Jongnam Kim\nApplied Sciences · MDPI · 2023'),('국문 초록 요약', '측면 영상에 OpenPose를 적용하고\n골격 각도로 정상·text neck·L자세 분류\n50명, 952장: 정확도 97.06%, F1 95.23%'),('활용 및 한계', '자세 지표와 분류 기준의 출발점\nPC 기반 보고 수치: Nano 성능 근거 아님\n개인화·연속 영상 평가는 추가 연구 대상')], takeaway='직접 관련 논문: 자세 분류 문제 설정. MDPI 저널 조건에 해당하며 최상위 학회 논문과 구분.', sources=['P1']),
 dict(section='APPENDIX · PAPER 02', title='Lite-HRNet: 고해상도 특징을 유지하는 경량 모델', sub='Lite-HRNet: A Lightweight High-Resolution Network', cards=[('Paper information', 'Changqian Yu et al.\nCVPR 2021 · pp. 10440–10450\n과제의 Top-tier 학회 조건 충족'),('국문 초록 요약', '고비용 1×1 합성곱을 줄이기 위해\n조건부 채널 가중치를 활용\n여러 해상도의 정보를 교환하며 자세 추정'),('활용 및 한계', '경량 키포인트 추정의 1차 구현 후보\n고정 사용자 영역으로 입력 구성 단순화\n측면·책상 가림 및 Nano 배포는 별도 검증')], takeaway='활용 방식: 원본 네트워크를 기준 모델로 사용하고 개인화·시간 조건의 효과를 분리 평가.', sources=['P2']),
 dict(section='APPENDIX · PAPER 03', title='Lite Pose: 엣지 장치를 위한 효율적 구조 설계', sub='Lite Pose: Efficient Architecture Design for 2D Human Pose Estimation', cards=[('Paper information', 'Yihan Wang, Muyang Li, Han Cai,\nWei-Ming Chen, Song Han\nCVPR 2022 · pp. 13126–13136'),('국문 초록 요약', '저연산 조건에서 고해상도 분기의\n중복을 분석하고 단일 분기 구조 설계\n특징 융합·큰 커널로 추정 성능 보완'),('활용 및 한계', '경량 구조 선택과 효율 평가의 비교 근거\n다중 인물 연구를 단일 사용자로 적용 검토\n논문상 모바일 성능을 Nano에 대입하지 않음')], takeaway='논문 3편의 역할: 자세 분류 근거(P1) + 경량 기준 모델(P2) + 효율 설계 비교(P3).', sources=['P3']),
 dict(section='APPENDIX · EVALUATION', title='감지부터 자세 복귀까지 검증', sub='측면 카메라 → 키포인트 → 기준 대비 변화 → 지속 시간 판단 → 알림', cards=[('데이터와 정답', '수집 동의를 받은 성인 15–20명 목표\n귀·어깨·엉덩이가 보이는 측면 영상\n사용자 단위 학습·검증·시험 분리'),('감지와 시스템 평가', 'Macro-F1 / 전방머리자세 재현율\n시간당 오경보 / 프레임 처리 지연 p95\n최대 전체 RAM / FPS / 판정 가능 비율'),('피드백 효과 평가', '알림 유·무 순서를 바꾼 비교 실험\n알림 후 10초 내 기준 자세 복귀율\n기준 이탈 지속 시간과 주관적 피로도')], takeaway='어깨 키포인트는 C7이 아니므로 귀–어깨 각도를 임상 CVA로 표기하지 않는다.'),
 dict(section='APPENDIX · ABSTRACT', title='연구 초록', sub='연구계획형 초록 · 아직 실험하지 않은 성능이나 치료 효과는 포함하지 않음', abstract=ABSTRACT, takeaway='키워드: 전방머리자세 · 경량 자세 추정 · Jetson Nano · 엣지 AI · 개인화 피드백'),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG='0D1728'; PANEL='16243A'; TEXT='F5F7FC'; MUTED='ADBBD0'; ACCENT='45D6BE'
def rgb(h): return RGBColor.from_string(h)
def box(slide,x,y,w,h,color,round=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background()
    return sh
def text(slide,x,y,w,h,content,size=20,color=TEXT,bold=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=Pt(2); tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name='Noto Sans CJK KR'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_after=Pt(9)
    return sh
def base(section,title,sub,n):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid(); sl.background.fill.fore_color.rgb=rgb(BG)
    box(sl,.55,.45,.08,.22,ACCENT)
    text(sl,.78,.39,11,.35,section,11,ACCENT,True)
    text(sl,.65,.97,12,1.05,title,28,TEXT,True)
    text(sl,.65,2.04,12,.76,sub,15,MUTED)
    text(sl,12.1,7.06,.6,.25,f'{n:02d}',10,MUTED)
    return sl

md=['# Jetson Nano 4GB 기반 개인화 자세 교정 지원 시스템','', '조사 기준일: 2026-09-13. 연구 제안 단계이며 하드웨어 실측과 사용자 실험은 수행하지 않았다.','']
for n,d in enumerate(slides,1):
    sl=base(d['section'],d['title'],d.get('sub',''),n)
    if d.get('cover'):
        for i,(head,body) in enumerate(d['cards']):
            text(sl,.7,3.25+i*1.07,10,.34,head,14,ACCENT,True)
            text(sl,.7,3.66+i*1.07,11,.5,body,20)
        box(sl,11.7,3.35,.2,2.3,ACCENT)
    elif 'abstract' in d:
        box(sl,.65,2.92,12.03,3.5,PANEL,True)
        text(sl,.96,3.13,11.38,3.05,d['abstract'],18)
    else:
        cards=d['cards']; cols=2 if len(cards)==4 else len(cards)
        w=(12.03-(cols-1)*.24)/cols
        for i,(head,body) in enumerate(cards):
            x=.65+(i%cols)*(w+.24); y=2.88+(i//cols)*1.75
            h=1.60 if len(cards)==4 else 3.25
            box(sl,x,y,w,h,PANEL,True)
            text(sl,x+.20,y+.16,w-.40,.43,head,17,ACCENT,True)
            body_shape=text(sl,x+.20,y+.70,w-.40,h-.77,body,15 if len(cards)==4 else 17)
            if len(cards)==4:
                for paragraph in body_shape.text_frame.paragraphs:
                    paragraph.space_after=Pt(3)
    text(sl,.65,6.5,12.03,.5,d.get('takeaway',''),13,ACCENT)
    refs=d.get('sources',[])
    if refs:
        tx=text(sl,.65,7.06,11.2,.27,'',9,MUTED)
        p=tx.text_frame.paragraphs[0]
        for k in refs:
            run=p.add_run(); run.text=f'[{k}] {S[k][0]}   '; run.hyperlink.address=S[k][1]
            run.font.name='Noto Sans CJK KR'; run.font.size=Pt(8); run.font.color.rgb=rgb(MUTED)
    notes='\n'.join([d['title'],d.get('sub',''),d.get('takeaway','')]+[f'[{k}] {S[k][0]}\n{S[k][1]}' for k in refs])
    sl.notes_slide.notes_text_frame.text=notes
    md += [f'## {n}. '+d['title'].replace('\n',' '),'',d.get('sub',''),'']
    if 'abstract' in d: md += [d['abstract'],'']
    for head,body in d.get('cards',[]): md += [f'**{head}**', '',body.replace('\n','  \n'),'']
    md += [d.get('takeaway',''),'']
    for k in refs: md += [f'[{k}: {S[k][0]}]({S[k][1]})','']

for group in [['P1','P2','P3','H1','H2'],['B1','B2','B3','M1','M2','M3']]:
    sl=base('APPENDIX · REFERENCES','참고문헌 및 조사 출처','각 출처명 클릭 시 원문으로 이동 · 열람일 2026.09.13',len(prs.slides)+1)
    for i,k in enumerate(group):
        y=2.92+i*.55
        sh=text(sl,.8,y,11.7,.45,f'[{k}] {S[k][0]}',16)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address=S[k][1]
    text(sl,.8,6.45,11.8,.5,'기업 발표는 상용화·사업 동향 근거로 사용. 제안 시스템의 효과·시장 규모와 구분.',13,ACCENT)
    sl.notes_slide.notes_text_frame.text='\n\n'.join(f'[{k}] {S[k][0]}\n{S[k][1]}' for k in group)

prs.save(OUT/'자세교정_연구제안.pptx')
(OUT/'논문조사_초록_발표원고.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'연구초록.txt').write_text('Jetson Nano 4GB 기반 개인화 자세 교정 지원 시스템\n\n'+ABSTRACT+'\n\n키워드: 전방머리자세, 경량 자세 추정, Jetson Nano, 엣지 AI, 개인화 피드백\n',encoding='utf-8')
print(f'Created {len(prs.slides)} slides in {OUT}')
